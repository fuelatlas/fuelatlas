"""Synthetic workbooks mirroring the Oil Bulletin layout.

Real fixtures would mean committing a 4.5 MB binary, so the tests rebuild the
shapes that matter instead: the date in A2, country names in column A, the
merged country-code column of the rate sheets, and the per-row gaps that make
rates carry forward.
"""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

PRODUCT_HEADERS = [
    "Euro-super 95  (I)",
    "Gas oil automobile Automotive gas oil Dieselkraftstoff",
    " Gas oil de chauffage Heating gas oil Heizoel",
    " Fuel oil - Schweres Heizöl (III) Soufre",
    " Fuel oil -Schweres Heizöl (III) Soufre ",
    "GPL pour moteur LPG motor fuel",
]

ALL_COUNTRIES = [
    "Austria", "Belgium", "Bulgaria", "Croatia", "Cyprus", "Czechia", "Denmark",
    "Estonia", "Finland", "France", "Germany", "Greece", "Hungary", "Ireland",
    "Italy", "Latvia", "Lithuania", "Luxembourg", "Malta", "Netherlands",
    "Poland", "Portugal", "Romania", "Slovakia", "Slovenia", "Spain", "Sweden",
]


def write_price_workbook(path, date: dt.date, values: dict[str, list], aggregates=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["in EUR", *PRODUCT_HEADERS])
    ws.append([date, "1000 l", "1000 l", "1000 l", "t", "t", "1000 l"])
    for country in ALL_COUNTRIES:
        ws.append([country, *values.get(country, [1000.0, 1200.0, None, None, None, None])])
    if aggregates:
        ws.append(["CE/EC/EG EUR27_2020 (IV)\nMoyenne pondéré", 2000.0, 2100.0, None, None, None, None])
        ws.append(["CE/EC/EG Euro Area 20 (V)\nMoyenne pondér", 2050.0, 2150.0, None, None, None, None])
    wb.save(path)
    return path


def write_rate_workbook(path, sheets: dict[str, list[tuple]]):
    """``sheets`` maps a sheet name to rows of ``(code_or_None, since, *values)``."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append([None, None, f"{name} header"])
        ws.append(["CTR"])
        ws.append([None, "Since:", *PRODUCT_HEADERS])
        ws.append([None, None, "1000 l", "1000 l", "1000 l", "t", "t", "1000 l"])
        for row in rows:
            code = f"{row[0]}_" if row[0] else None
            ws.append([code, *row[1:]])
    wb.save(path)
    return path


@pytest.fixture
def week() -> dt.date:
    return dt.date(2026, 9, 14)
